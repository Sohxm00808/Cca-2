"""
Name : Soham Ashok Jagtap
PRN  : 1302250127
Creates input_data.xlsx (sheets: Training, Input) used by traditional_vs_ml.py
Run  : python create_input_excel.py
"""
import random
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

COLS = ["ApplicantID", "MonthlyIncome", "CreditScore", "ExistingEMI", "LoanAmount"]


def make_row(i):
    inc = random.randrange(20000, 150000, 5000)
    cs = random.randrange(550, 850, 10)
    emi = int(inc * random.uniform(0.0, 0.7) // 500 * 500)
    loan = int(inc * random.uniform(5, 40) // 10000 * 10000)
    return [i, inc, cs, emi, loan]


def historical_label(r):
    # Past bank decisions: similar to the rules but slightly different, so ML differs from rules
    _, inc, cs, emi, loan = r
    score = (cs - 600) / 250 - emi / inc * 0.9 - loan / (inc * 60)
    return "Yes" if score > 0.05 else "No"


def create(file="input_data.xlsx"):
    random.seed(127)
    wb = Workbook()
    train = wb.active
    train.title = "Training"
    train.append(COLS + ["Approved"])
    for i in range(1, 61):
        r = make_row(i)
        train.append(r + [historical_label(r)])

    new = wb.create_sheet("Input")
    new.append(COLS)
    for i in range(101, 116):
        new.append(make_row(i))

    for sheet in (train, new):
        for c in sheet[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
        for col in "ABCDEF":
            sheet.column_dimensions[col].width = 16
    wb.save(file)
    print(file, "created")


if __name__ == "__main__":
    create()
